from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "experiments" / "ll_depth_p3_evidence_20260810" / "outputs"
REPORT_PATH = ROOT / "daily_report" / "赵浩骅-0810-LL作为DepthP3输入实验日报.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.first_child_found_in("w:tblCellMar")
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_in):
    for row in table.rows:
        for cell, width in zip(row.cells, widths_in):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_font(run, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, size=11, bold=False, color=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 12}.get(level, 11), bold=True, color="2E74B5" if level < 3 else "1F4D78")
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_font(r, size=11)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=9, color="555555")
    return p


def add_picture(doc, path: Path, width_in: float, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    add_caption(doc, caption)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for level, size, before, after, color in [
        (1, 16, 16, 8, "2E74B5"),
        (2, 13, 12, 6, "2E74B5"),
        (3, 12, 8, 4, "1F4D78"),
    ]:
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("S8 Depth P3 LL 输入合理性实验日报")
    set_font(r, size=9, color="666666")


def add_metrics_table(doc, summary):
    rows = [
        ("LL-depth Pearson r", f"{summary['ll_depth_corr']['mean']:.4f}", f"{summary['ll_depth_corr']['median']:.4f}", "LL 与原始 depth 的结构一致性，高说明 LL 保留主体深度几何。"),
        ("LL-depth PSNR", f"{summary['ll_depth_psnr']['mean']:.2f} dB", f"{summary['ll_depth_psnr']['median']:.2f} dB", "LL 作为平滑重构的保真度，反映低频分量并非丢弃 depth 信息。"),
        ("LL energy share", f"{summary['ll_energy_share']['mean']:.4f}", f"{summary['ll_energy_share']['median']:.4f}", "Haar 系数能量中由低频承载的比例，说明主体结构主要在 LL。"),
        ("HF-boundary AUC", f"{summary['hf_boundary_auc']['mean']:.4f}", f"{summary['hf_boundary_auc']['median']:.4f}", "高频幅值对 GT mask 边界的区分度，反映其边缘增强价值。"),
        ("LL-gradient boundary AUC", f"{summary['ll_grad_boundary_auc']['mean']:.4f}", f"{summary['ll_grad_boundary_auc']['median']:.4f}", "LL 平滑后仍保留的边界梯度信号，说明 LL 没有抹掉主体轮廓。"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(["指标", "Mean", "Median", "解释"]):
        hdr[i].text = text
        set_cell_shading(hdr[i], "F2F4F7")
    for metric, mean, median, meaning in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, [metric, mean, median, meaning]):
            cell.text = text
    set_cell_margins(table)
    set_table_width(table, [1.65, 1.0, 1.0, 2.85])
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_font(run, size=9.5)


def main() -> None:
    summary = json.loads((OUT_DIR / "summary.json").read_text(encoding="utf-8"))
    case_images = sorted((OUT_DIR / "case_visualizations").glob("*.png"))
    case_image = case_images[2] if len(case_images) >= 3 else case_images[0]

    doc = Document()
    style_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("工作日报")
    set_font(r, size=24, bold=True, color="0B2545")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("S8 模型 Depth P3 中 LL 作为 CoordAttV2 输入的合理性验证")
    set_font(r, size=13, color="1F4D78")

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    meta_data = [
        ("日期", "2026 年 8 月 10 日"),
        ("姓名", "赵浩骅"),
        ("实验目录", str(ROOT / "experiments" / "ll_depth_p3_evidence_20260810")),
        ("实验环境", "conda 环境 genye-yolo；数据集 Dataset/xinjiang_1500 val；样本数 200"),
    ]
    for row, (k, v) in zip(meta.rows, meta_data):
        row.cells[0].text = k
        row.cells[1].text = v
        set_cell_shading(row.cells[0], "F2F4F7")
    set_cell_margins(meta)
    set_table_width(meta, [1.35, 5.15])

    add_heading(doc, "一、今日工作概述", 1)
    add_para(doc, "今日围绕 S8 模型中 P3 wavelet-guided CoordAttV2 的关键问题开展验证：Depth P3 经 Haar 分解后，低频 LL 是否适合作为 CoordAttV2 的 depth 主输入，高频 LH/HL/HH 是否更适合作为 edge_gate 的边缘增强信号。")
    add_bullet(doc, "核验了当前 S8 代码路径：Depth P3 -> Haar split -> LL 进入 low_proj 与 CoordAttV2；LH/HL/HH 经高频分支生成 edge_gate。")
    add_bullet(doc, "在 experiments 目录下新建实验并完成可复现实验脚本，输出 CSV、JSON、结构图、统计图和单样本可视化图。")
    add_bullet(doc, "基于 200 张验证集图像统计 LL 与原始 depth 的结构一致性、低频能量占比、高频与 GT 边界的相关性。")

    add_heading(doc, "二、实验问题与判断标准", 1)
    add_para(doc, "本实验要回答的问题不是“wavelet 一定提升 AP”，而是“当前结构中 LL 作为 CoordAttV2 的 depth 输入是否有信号层面的合理性”。因此判断标准分成三层：")
    add_bullet(doc, "结构保真：LL 应与原始 depth 保持较高相关，说明它仍携带主要几何和主体区域信息。")
    add_bullet(doc, "频率分工：低频应承载大部分能量，高频应主要反映突变、边缘和局部细节。")
    add_bullet(doc, "边界价值：高频分量对 GT mask 边界应有正向区分度，适合做边缘增强门控，而不是直接替代主 depth 输入。")

    add_heading(doc, "三、S8 P3 wavelet-guided CoordAttV2 结构可视化", 1)
    add_picture(doc, OUT_DIR / "s8_p3_wavelet_flow.png", 6.25, "图 1：S8 P3 wavelet-guided CoordAttV2 数据流。LL 作为稳定几何输入，高频分量生成 edge_gate。")
    add_para(doc, "从结构上看，S8 将深度特征分成两类用途：LL 进入 CoordAttV2 主融合路径，承担深度主体结构和坐标方向信息；LH/HL/HH 被压缩为单通道 edge_gate，对 base 融合结果进行乘性增强。这样的设计避免把高频噪声直接混入主融合路径，同时保留边界增强能力。")

    add_heading(doc, "四、定量统计结果", 1)
    add_metrics_table(doc, summary)
    add_para(doc, "最核心的两个数字是 LL-depth Pearson r = 0.9572，以及 LL energy share = 0.9860。前者说明 LL 与原始深度图的空间结构高度一致；后者说明 Haar 分解后绝大多数能量位于低频，主体几何主要由 LL 承载。HF-boundary AUC 为 0.6299，说明高频对 GT 边界有中等正向区分能力，但不是完美边界检测器，这一点符合真实深度图中存在背景深度突变、传送带边缘和噪声点的实际情况。")
    add_picture(doc, OUT_DIR / "aggregate_charts.png", 6.3, "图 2：总体统计图。左：LL 与 depth 相关性；中：低频能量占比；右：边界对齐 AUC。")

    add_heading(doc, "五、单样本可视化解释", 1)
    add_picture(doc, case_image, 6.35, "图 3：验证集样本的 RGB、Depth、LL、HF magnitude、GT mask 和 |Depth-LL| 对比。")
    add_para(doc, "图 3 可以直观看到：Depth 与 LL upsampled 在主体深度区域和主要轮廓上高度一致，说明 LL 并不是把深度信息丢弃，而是对 depth 做了低通平滑；HF magnitude 主要亮在深度突变、物体边缘、传送带边界和孤立噪声点上，更像局部变化提示。")
    add_para(doc, "因此，把 LL 送入 CoordAttV2 可以让融合模块看到更稳定的深度几何；把 HF 单独作为 edge_gate，则可以让模型在边缘位置放大融合特征，而不会让高频噪声主导 RGB-D 主融合。")

    add_heading(doc, "六、结论", 1)
    add_para(doc, "当前证据支持“LL 作为 Depth P3 输入是合适的”。对 CoordAttV2 这类依赖全局池化和坐标方向建模的融合模块而言，输入应尽量稳定、结构清晰、语义连续。LL 正好满足这一点：它保留 depth 主体几何，又降低局部噪声和细碎高频的干扰。")
    add_para(doc, "高频分量 LH/HL/HH 更适合作为边缘增强分支，因为高频的优势在于捕捉突变和边界，但它同时包含背景边缘、传感器噪声和非目标深度变化。如果直接把高频作为 CoordAttV2 的主 depth 输入，可能会降低融合路径的稳定性。")

    add_heading(doc, "七、原因分析", 1)
    add_heading(doc, "1. 为什么 LL 更适合作为主 depth 输入", 2)
    add_bullet(doc, "CoordAttV2 通过水平和垂直方向的全局池化建模坐标注意力，主输入需要表达连续、稳定的空间结构。LL 是低频分量，天然保留主体几何、区域深度和整体轮廓。")
    add_bullet(doc, "实验中 LL 与原始 depth 的相关均值达到 0.9572，说明 LL 不是弱化 depth，而是把 depth 变成更平滑、更利于融合的几何先验。")
    add_bullet(doc, "低频能量占比 0.9860，说明大部分深度结构能量本来就在 LL 中，主融合路径使用 LL 不会造成主体信息损失。")

    add_heading(doc, "2. 为什么高频更适合做 edge_gate", 2)
    add_bullet(doc, "LH/HL/HH 对局部突变敏感，能强调物体边缘、遮挡边界和深度断裂，因此适合提供边缘增强。")
    add_bullet(doc, "HF-boundary AUC 为 0.6299，说明高频确实与标注边界存在正相关；但该值不是极高，说明高频还包含背景和噪声，所以作为辅助门控更稳妥。")
    add_bullet(doc, "edge_gate 是乘性增强分支，作用是调节 base 特征，而不是单独决定融合结果，这种设计可以控制高频噪声带来的风险。")

    add_heading(doc, "3. 为什么该实验仍需 AP 消融补强", 2)
    add_bullet(doc, "本次实验是数据级和信号级分析，证明当前设计在理论和可视化上合理，但不能直接等价于最终 mAP 提升。")
    add_bullet(doc, "若要形成论文级闭环，后续应在相同训练配置下比较：LL as CoordAtt input、raw depth as CoordAtt input + HF gate、HF/mixed depth as CoordAtt input。")
    add_bullet(doc, "如果三组训练结果显示 LL 输入版本 AP 更优或更稳定，则可以把本次统计和可视化作为机制解释，把 AP 消融作为最终性能证据。")

    add_heading(doc, "八、后续计划", 1)
    add_bullet(doc, "整理 ablation 方案，优先实现 raw depth 输入和 mixed depth 输入两个对照分支。")
    add_bullet(doc, "在同一数据集、同一训练轮数、同一随机种子范围内比较 mask mAP、box mAP 和类别级指标。")
    add_bullet(doc, "将本次可视化图和后续 AP 消融结果合并进 S8 网络结构说明或论文实验章节。")

    doc.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    main()
